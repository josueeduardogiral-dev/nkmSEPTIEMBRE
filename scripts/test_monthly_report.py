from datetime import date, time
import os
from pathlib import Path
import sys
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

os.environ.setdefault(
    'DATABASE_URL',
    'postgresql://test:test@example.invalid:5432/postgres',
)
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'control_asistencia.settings')

import django
django.setup()

import openpyxl
from io import BytesIO
from unittest.mock import patch
from django.test import RequestFactory
from django.urls import resolve, reverse
from app.services import ReporteService
from app.views import (
    _nombre_hoja_empleado,
    exportar_reporte_mensual_empleados,
    pagina_descarga_excel,
)


class MonthlyReportTests(unittest.TestCase):
    def test_august_includes_fifth_week_and_day_31(self):
        ranges = ReporteService.rangos_semanales_mes(date(2026, 8, 31))
        self.assertEqual(
            ranges,
            [
                (1, date(2026, 8, 3), date(2026, 8, 7)),
                (2, date(2026, 8, 10), date(2026, 8, 14)),
                (3, date(2026, 8, 17), date(2026, 8, 21)),
                (4, date(2026, 8, 24), date(2026, 8, 28)),
                (5, date(2026, 8, 31), date(2026, 8, 31)),
            ],
        )

    def test_partial_weeks_stay_inside_selected_month(self):
        ranges = ReporteService.rangos_semanales_mes(date(2024, 2, 10))
        self.assertEqual(ranges[0], (1, date(2024, 2, 1), date(2024, 2, 2)))
        self.assertEqual(ranges[-1], (5, date(2024, 2, 26), date(2024, 2, 29)))

    def test_excel_sheet_names_are_valid_and_unique(self):
        empleado = type('Empleado', (), {
            'nombre_completo': 'Nombre / muy largo : para [Excel] empleado repetido',
            'pk': 1,
        })()
        usados = set()
        primero = _nombre_hoja_empleado(empleado, usados)
        segundo = _nombre_hoja_empleado(empleado, usados)
        self.assertLessEqual(len(primero), 31)
        self.assertNotEqual(primero, segundo)
        self.assertFalse(any(char in primero for char in '[]:*?/\\'))

    @patch('app.views.RegistroAsistencia.objects')
    @patch('app.views.Empleado.objects')
    def test_empty_database_still_generates_valid_workbook(self, empleados, registros):
        empleados.order_by.return_value = []
        registros.filter.return_value.select_related.return_value.order_by.return_value = []
        undecorated = exportar_reporte_mensual_empleados.__wrapped__
        response = undecorated(object())
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Sin empleados'])
        self.assertIn('No hay empleados', workbook.active['A1'].value)

    @patch('app.views.RegistroAsistencia.objects')
    @patch('app.views.Empleado.objects')
    def test_selected_month_controls_filter_weeks_and_filename(self, empleados, registros):
        empleado = type('Empleado', (), {
            'id_empleado': 7,
            'nombre_completo': 'Empleado de prueba',
            'pk': 7,
        })()
        registro_31 = type('Registro', (), {
            'empleado_id': 7,
            'fecha_registro': date(2026, 8, 31),
            'hora_registro': time(9, 0),
            'tipo': type('Tipo', (), {'nombre_asistencia': 'Entrada'})(),
        })()
        empleados.order_by.return_value = [empleado]
        queryset = registros.filter.return_value.select_related.return_value.order_by.return_value
        queryset.__iter__.return_value = iter([registro_31])
        request = type('Request', (), {'GET': {'mes': '2026-08'}})()

        response = exportar_reporte_mensual_empleados.__wrapped__(request)
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        worksheet = workbook['Empleado de prueba']
        week_labels = [
            cell.value for cell in worksheet['A']
            if isinstance(cell.value, str) and cell.value.startswith('Semana ')
        ]

        registros.filter.assert_called_once_with(
            fecha_registro__range=(date(2026, 8, 1), date(2026, 8, 31))
        )
        self.assertEqual(len(week_labels), 5)
        self.assertEqual(week_labels[0], 'Semana 1 (03/08/2026 - 07/08/2026)')
        self.assertEqual(week_labels[-1], 'Semana 5 (31/08/2026 - 31/08/2026)')
        values = list(worksheet.values)
        self.assertTrue(any(row[0] == '31/08/2026' and row[1] == '09:00' for row in values))
        self.assertIn('reporte_mensual_empleados_2026_08.xlsx', response['Content-Disposition'])

    def test_invalid_month_is_rejected(self):
        request = type('Request', (), {'GET': {'mes': 'agosto-2026'}})()
        response = exportar_reporte_mensual_empleados.__wrapped__(request)
        self.assertEqual(response.status_code, 400)
        self.assertIn('Mes inválido', response.content.decode('utf-8'))

    def test_month_outside_2026_is_rejected(self):
        request = type('Request', (), {'GET': {'mes': '2025-08'}})()
        response = exportar_reporte_mensual_empleados.__wrapped__(request)
        self.assertEqual(response.status_code, 400)
        self.assertIn('año 2026', response.content.decode('utf-8'))

    def test_existing_attendance_download_uses_monthly_report(self):
        url = reverse('descargar_excel')
        self.assertEqual(url, '/login/descargar/asistencia')
        self.assertIs(resolve(url).func, exportar_reporte_mensual_empleados)

    @patch('app.views.timezone.localdate', return_value=date(2026, 9, 7))
    def test_download_page_first_action_contains_month_selector(self, _localdate):
        request = RequestFactory().get(reverse('pagina_descarga_excel'))
        response = pagina_descarga_excel.__wrapped__(request)
        html = response.content.decode('utf-8')
        self.assertIn('name="mes"', html)
        self.assertIn('value="2026-09"', html)
        self.assertIn('todas las semanas', html)
        self.assertIn('min="2026-01"', html)
        self.assertIn('max="2026-12"', html)
        self.assertIn(reverse('descargar_excel'), html)


if __name__ == '__main__':
    unittest.main()
